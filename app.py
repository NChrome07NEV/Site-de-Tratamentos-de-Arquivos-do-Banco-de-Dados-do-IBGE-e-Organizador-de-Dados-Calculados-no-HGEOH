import streamlit as st
import pandas as pd
import io
import openpyxl
from openpyxl.styles import Font, Alignment

# Configuração da página sem emojis
st.set_page_config(
    page_title="Organizador e Formator de Arquivos Brutos das Estações Gravimétricas Gerados no Site do Banco de Dados do IBGE ",
    layout="wide"
)

st.title("Acesso Restrito - Organizador e Formatador de Arquivos Brutos das Estações Gravimétricas Gerados no Site do Banco de Dados do IBGE (RJ)")

# Função para formatar o arquivo Excel com OpenPyXL
def aplicar_estilo_excel(df, nome_aba='Sheet1'):
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, sheet_name=nome_aba, index=False)
        workbook = writer.book
        worksheet = writer.sheets[nome_aba]
        
        # Estilo solicitado: Times New Roman, Tamanho 12, Centralizado
        fonte_customizada = Font(name='Times New Roman', size=12)
        alinhamento_central = Alignment(horizontal='center', vertical='center')
        
        for row in worksheet.iter_rows(min_row=1, max_row=worksheet.max_row, min_col=1, max_col=worksheet.max_column):
            for cell in row:
                cell.font = fonte_customizada
                cell.alignment = alinhamento_central
                
        # Auto-ajuste de largura das colunas
        for col in worksheet.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            worksheet.column_dimensions[col_letter].width = max(max_len + 3, 12)
            
    return output.getvalue()

# Função para ordenar dataframe por data
def ordenar_por_data(df):
    colunas_data = [col for col in ['VISITA', 'MEDICAOPLAN', 'CALCULOGRA', 'MEDICAOALT'] if col in df.columns]
    if colunas_data:
        col_ref = colunas_data[0]
        df_temp = df.copy()
        df_temp['_data_sort'] = pd.to_datetime(df_temp[col_ref], format='%d/%m/%Y', errors='coerce')
        df_sorted = df_temp.sort_values(by='_data_sort').drop(columns=['_data_sort'])
        return df_sorted
    return df

# Função para gerar o TXT estrito com 3 colunas (ESTACAO, LATITUDE, LONGITUDE)
def gerar_txt_estacoes(df):
    df_txt = df[['ESTACAO', 'LATITUDE', 'LONGITUDE']].copy()
    df_txt = df_txt.dropna(subset=['ESTACAO', 'LATITUDE', 'LONGITUDE'])
    return df_txt.to_csv(index=False, header=False, sep=' ')

# --- ETAPA 1: Organização dos Dados do IBGE ---
st.header("Etapa 1: Organização dos Dados do IBGE")

arquivo_etapa1 = st.file_uploader("Selecione o arquivo Excel bruto da Etapa 1 (.xlsx)", type=["xlsx"], key="up1")

if arquivo_etapa1 is not None:
    try:
        df_etapa1 = pd.read_excel(arquivo_etapa1)
        
        # Ordenar por data e formatar
        df_etapa1_proc = ordenar_por_data(df_etapa1)
        
        excel_bytes = aplicar_estilo_excel(df_etapa1_proc, 'Sheet1')
        txt_bytes = gerar_txt_estacoes(df_etapa1_proc)
        
        st.success("Dados da Etapa 1 processados com sucesso!")
        
        col1, col2 = st.columns(2)
        with col1:
            st.download_button(
                label="Baixar Excel Organizado (Etapa 1)",
                data=excel_bytes,
                file_name="00_IBGE_Estacoes_Tratadas_Organizadas.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
        with col2:
            st.download_button(
                label="Baixar TXT de Entrada (Etapa 1)",
                data=txt_bytes,
                file_name="Estacoes_Entrada_HGEOHNOR2020.txt",
                mime="text/plain"
            )
    except Exception as e:
        st.error(f"Erro ao processar o arquivo da Etapa 1: {e}")

st.markdown("---")

# --- ETAPA 2: Organização dos Dados do HGEOH ---
st.header("Etapa 2: Organização dos Dados do HGEOH")

arquivo_etapa2 = st.file_uploader("Selecione o arquivo Excel da Etapa 2 (.xlsx)", type=["xlsx"], key="up2")

if arquivo_etapa2 is not None:
    try:
        df_etapa2 = pd.read_excel(arquivo_etapa2)
        
        # Ordenar por data e formatar com o mesmo padrão Times New Roman 12 centralizado
        df_etapa2_proc = ordenar_por_data(df_etapa2)
        excel_bytes_etapa2 = aplicar_estilo_excel(df_etapa2_proc, 'Sheet1')
        
        st.success("Dados da Etapa 2 processados com sucesso!")
        
        st.download_button(
            label="Baixar Excel Organizado e Formatado (Etapa 2)",
            data=excel_bytes_etapa2,
            file_name="Dados Organizados - EG.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
    except Exception as e:
        st.error(f"Erro ao processar o arquivo da Etapa 2: {e}")
